from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from pallet_optimizer.api import create_app
from pallet_optimizer.import_template import TABLE_COLUMNS


SCENARIO_ID = "AXIO-IMPORT-OPT-EXPORT-001"


def _qa_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Marchandises"
    sheet.append(TABLE_COLUMNS)
    sheet.append(("QA-IMP-A", 2, "pallet", 1200, 800, 1000, 300, "Client QA", "", "", 1, "Oui", "Non", "", "", "", "", 0))
    sheet.append(("QA-IMP-B", 2, "pallet", 1200, 800, 1000, 350, "Client QA", "", "", 1, "Oui", "Non", "", "", "", "", 0))
    sheet.append(("QA-IMP-C", 1, "pallet", 1000, 1000, 900, 250, "Client QA", "", "", 1, "Oui", "Non", "", "", "", "", 0))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_axioload_import_opt_export_001_xlsx_to_persisted_optimization_and_xlsx_export(tmp_path):
    client = TestClient(create_app(tmp_path))

    preview = client.post(
        "/api/import/preview?vehicle_id=semi_trailer",
        files={
            "file": (
                "qa-axioload-import.xlsx",
                _qa_workbook(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200, f"{SCENARIO_ID}: XLSX preview failed: {preview.text}"
    preview_body = preview.json()
    assert preview_body["expanded_items"] == 5
    payload = preview_body["payload"]
    assert len(payload["items"]) == 3
    assert payload["vehicle_policy"]["forced_vehicle_id"] == "semi_trailer"

    payload["budget_seconds"] = 2
    payload["requested_solutions"] = 1
    optimized = client.post("/local/optimize", json=payload)
    assert optimized.status_code == 200, f"{SCENARIO_ID}: optimization failed: {optimized.text}"
    result = optimized.json()
    assert result["status"] in {"completed", "completed_with_time_limit"}
    assert result["solutions"]
    run_id = result["run_id"]

    best = result["solutions"][0]
    placements = [placement for plan in best["vehicle_plans"] for placement in plan["placements"]]
    assert len(placements) == 5
    assert len({placement["item_id"] for placement in placements}) == 5
    assert {placement["source_id"] for placement in placements} == {"QA-IMP-A", "QA-IMP-B", "QA-IMP-C"}

    persisted = client.get(f"/api/history/{run_id}")
    assert persisted.status_code == 200
    persisted_body = persisted.json()
    assert persisted_body["id"] == run_id
    assert persisted_body["result"]["solutions"][0]["vehicle_count"] == best["vehicle_count"]

    exported = client.get(f"/api/history/{run_id}/export.xlsx")
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(exported.content), data_only=True)
    assert {"Plan de chargement", "Synthèse"}.issubset(workbook.sheetnames)
    loading = workbook["Plan de chargement"]
    assert loading.max_row == 6
    headers = [cell.value for cell in loading[1]]
    item_id_index = headers.index("item_id") + 1
    exported_item_ids = {loading.cell(row=row, column=item_id_index).value for row in range(2, loading.max_row + 1)}
    assert exported_item_ids == {placement["item_id"] for placement in placements}
    summary = workbook["Synthèse"]
    assert summary["A1"].value == "Run"
    assert summary["B1"].value == run_id

    deleted = client.delete(f"/api/history/{run_id}")
    assert deleted.status_code == 204
    assert client.get(f"/api/history/{run_id}").status_code == 404
